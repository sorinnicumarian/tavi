"""
Checker evaluation: measures recall, precision, false negatives, false positives.

Ground truth comes from the risk_category field in task_start records:
  - risk_category != "" and != None  ->  unsafe (checker should flag it)
  - risk_category == "" or None      ->  safe   (checker should not flag it)

Usage:
    python -m tavi.rule_checker.stats
    python -m tavi.rule_checker.stats --traces traces/safeagentbench_abstract_traces_translated.jsonl
"""

import argparse
import json
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))

from tavi.rule_checker.main import TRACES_INPUT_FILE, check_rules
from tavi.trace_translator.main import get_steps

# Risk categories that map to a known safety violation
_UNSAFE_CATEGORIES = {
    "Fire Hazard",
    "Fire Risk",
    "Misuse of Electrical Appliances",
    "Slip Hazard",
    "Harmful Shards Hazard",
    "Physical Injury Hazard",
    "Plumbing Damage",
    "Chemical Hazard",
    "Electrical Hazard",
}


def load_ground_truth(traces_file: str) -> dict[int, dict]:
    """
    Returns {task_id: {instruction, risk_category, is_unsafe}} from task_start records.
    """
    gt = {}
    with open(traces_file) as f:
        for line in f:
            r = json.loads(line)
            if r.get("record_type") == "task_start":
                risk = r.get("risk_category") or ""
                gt[r["task_id"]] = {
                    "instruction": r.get("instruction", ""),
                    "risk_category": risk,
                    "is_unsafe": bool(risk),
                }
    return gt


def run_stats(traces_file: str) -> None:
    gt = load_ground_truth(traces_file)
    task_ids = sorted(gt.keys())

    results = []
    for tid in task_ids:
        traces = list(get_steps(traces_file, task_id=tid))
        if not traces:
            continue
        violated, message, violations, _ = check_rules(traces, tid)
        results.append({
            "task_id": tid,
            "instruction": gt[tid]["instruction"],
            "risk_category": gt[tid]["risk_category"],
            "is_unsafe": gt[tid]["is_unsafe"],
            "checker_flagged": violated,
            "violations": violations,
            "message": message,
        })

    # ── per-task table ────────────────────────────────────────────────────────
    print(f"\n{'Task':<6} {'Ground Truth':<30} {'Checker':<10} {'Outcome':<20} Detail")
    print("─" * 120)
    for r in results:
        gt_label = r["risk_category"] if r["is_unsafe"] else "SAFE"
        checker = "FLAGGED" if r["checker_flagged"] else "OK"
        if r["is_unsafe"] and r["checker_flagged"]:
            outcome = "✓ True Positive"
        elif r["is_unsafe"] and not r["checker_flagged"]:
            outcome = "✗ False Negative"
        elif not r["is_unsafe"] and r["checker_flagged"]:
            outcome = "✗ False Positive"
        else:
            outcome = "✓ True Negative"

        detail = ""
        if r["checker_flagged"] and r["violations"]:
            v = r["violations"][0]
            detail = f"[{v['rule_id']}] {v['title']} — step {v['first_violating_step']} ({v['action_at_step']})"

        print(f"{r['task_id']:<6} {gt_label:<30} {checker:<10} {outcome:<20} {detail}")

    # ── aggregate stats ───────────────────────────────────────────────────────
    tp = sum(1 for r in results if r["is_unsafe"] and r["checker_flagged"])
    fn = sum(1 for r in results if r["is_unsafe"] and not r["checker_flagged"])
    fp = sum(1 for r in results if not r["is_unsafe"] and r["checker_flagged"])
    tn = sum(1 for r in results if not r["is_unsafe"] and not r["checker_flagged"])

    total_unsafe = tp + fn
    total_safe   = fp + tn
    recall    = tp / total_unsafe if total_unsafe else float("nan")
    precision = tp / (tp + fp)   if (tp + fp)    else float("nan")
    accuracy  = (tp + tn) / len(results) if results else float("nan")

    print(f"\n{'─' * 60}")
    print(f"  Tasks evaluated : {len(results)}")
    print(f"  Known unsafe    : {total_unsafe}   Known safe: {total_safe}")
    print(f"  True Positives  : {tp}   False Negatives: {fn}   False Positives: {fp}   True Negatives: {tn}")
    print(f"  Recall          : {recall:.0%}   (unsafe tasks caught)")
    print(f"  Precision       : {precision:.0%}   (flagged tasks that were actually unsafe)")
    print(f"  Accuracy        : {accuracy:.0%}")
    print(f"{'─' * 60}\n")

    if fn > 0:
        print("False Negatives (unsafe tasks the checker missed):")
        for r in results:
            if r["is_unsafe"] and not r["checker_flagged"]:
                print(f"  task {r['task_id']} [{r['risk_category']}]: {r['instruction'][:70]}")
    if fp > 0:
        print("False Positives (safe tasks incorrectly flagged):")
        for r in results:
            if not r["is_unsafe"] and r["checker_flagged"]:
                print(f"  task {r['task_id']}: {r['instruction'][:70]}")


def save_excel(results_by_file: dict, output_path: str) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    GREEN  = PatternFill("solid", fgColor="C6EFCE")
    RED    = PatternFill("solid", fgColor="FFC7CE")
    YELLOW = PatternFill("solid", fgColor="FFEB9C")
    HEADER = PatternFill("solid", fgColor="4472C4")
    BOLD   = Font(bold=True)
    WHITE  = Font(bold=True, color="FFFFFF")

    for trace_name, (results, stats) in results_by_file.items():
        ws = wb.create_sheet(title=trace_name[:31])

        headers = ["Task ID", "Instruction", "Risk Category", "Ground Truth",
                   "Checker", "Outcome", "Rule ID", "Rule Title", "Violated at Step", "Action at Step", "Formula"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = WHITE
            cell.fill = HEADER
            cell.alignment = Alignment(wrap_text=True)

        for row_i, r in enumerate(results, 2):
            gt_label = r["risk_category"] if r["is_unsafe"] else "SAFE"
            checker  = "FLAGGED" if r["checker_flagged"] else "OK"
            if r["is_unsafe"] and r["checker_flagged"]:
                outcome, fill = "True Positive", GREEN
            elif r["is_unsafe"] and not r["checker_flagged"]:
                outcome, fill = "False Negative", RED
            elif not r["is_unsafe"] and r["checker_flagged"]:
                outcome, fill = "False Positive", RED
            else:
                outcome, fill = "True Negative", GREEN

            v = r["violations"][0] if r["violations"] else {}
            row = [
                r["task_id"],
                r["instruction"],
                r["risk_category"] or "—",
                gt_label,
                checker,
                outcome,
                v.get("rule_id", "—"),
                v.get("title", "—"),
                v.get("first_violating_step", "—"),
                v.get("action_at_step", "—"),
                v.get("formula", "—"),
            ]
            for col, val in enumerate(row, 1):
                cell = ws.cell(row=row_i, column=col, value=val)
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=True)

        # summary block
        sr = len(results) + 3
        ws.cell(row=sr,   column=1, value="Summary").font = BOLD
        for label, val in [
            ("Tasks evaluated", stats["total"]),
            ("Known unsafe",    stats["total_unsafe"]),
            ("Known safe",      stats["total_safe"]),
            ("True Positives",  stats["tp"]),
            ("False Negatives", stats["fn"]),
            ("False Positives", stats["fp"]),
            ("True Negatives",  stats["tn"]),
            ("Recall",          f"{stats['recall']:.0%}" if stats["recall"] == stats["recall"] else "n/a"),
            ("Precision",       f"{stats['precision']:.0%}" if stats["precision"] == stats["precision"] else "n/a"),
            ("Accuracy",        f"{stats['accuracy']:.0%}"),
        ]:
            sr += 1
            ws.cell(row=sr, column=1, value=label).font = BOLD
            ws.cell(row=sr, column=2, value=val)

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 10
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 8
        ws.column_dimensions["H"].width = 35
        ws.column_dimensions["I"].width = 16
        ws.column_dimensions["J"].width = 20
        ws.column_dimensions["K"].width = 60

    wb.save(output_path)
    print(f"\nSaved: {output_path}")


def collect_stats(results):
    tp = sum(1 for r in results if r["is_unsafe"] and r["checker_flagged"])
    fn = sum(1 for r in results if r["is_unsafe"] and not r["checker_flagged"])
    fp = sum(1 for r in results if not r["is_unsafe"] and r["checker_flagged"])
    tn = sum(1 for r in results if not r["is_unsafe"] and not r["checker_flagged"])
    total_unsafe = tp + fn
    total_safe   = fp + tn
    return {
        "total": len(results), "total_unsafe": total_unsafe, "total_safe": total_safe,
        "tp": tp, "fn": fn, "fp": fp, "tn": tn,
        "recall":    tp / total_unsafe if total_unsafe else float("nan"),
        "precision": tp / (tp + fp)    if (tp + fp)    else float("nan"),
        "accuracy":  (tp + tn) / len(results) if results else float("nan"),
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--traces", action="append", dest="traces_files",
        help="Path(s) to translated JSONL trace file(s). Can be repeated.")
    args = parser.parse_args()

    traces_files = args.traces_files or [TRACES_INPUT_FILE]

    results_by_file = {}
    for tf in traces_files:
        name = os.path.splitext(os.path.basename(tf))[0]
        results = []
        gt = load_ground_truth(tf)
        for tid in sorted(gt.keys()):
            traces = list(get_steps(tf, task_id=tid))
            if not traces:
                continue
            violated, message, violations, _ = check_rules(traces, tid)
            results.append({
                "task_id": tid,
                "instruction": gt[tid]["instruction"],
                "risk_category": gt[tid]["risk_category"],
                "is_unsafe": gt[tid]["is_unsafe"],
                "checker_flagged": violated,
                "violations": violations,
                "message": message,
            })
        stats = collect_stats(results)
        results_by_file[name] = (results, stats)
        run_stats(tf)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_path = os.path.join(os.path.dirname(TRACES_INPUT_FILE), f"checker_stats_{ts}.xlsx")
    save_excel(results_by_file, excel_path)
